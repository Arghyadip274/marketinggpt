"""Document ingestion and embedding pipeline.

Integrates with LangChain text splitters and ChromaDB (or memory store fallback)
for processing raw markdown, PDF, DOCX, and TXT files into vector embeddings.
"""

import logging
import os
from pathlib import Path
from typing import Any, Dict, List, Optional

from app.config.settings import settings
from app.rag.retriever import KnowledgeRetriever

logger = logging.getLogger(__name__)


class IngestError(Exception):
    """Base exception for document ingestion failures."""


class DocumentIngester:
    """Handles the ingestion of raw documents into the vector store."""

    def __init__(self, documents_dir: str | None = None, retriever: KnowledgeRetriever | None = None):
        """Initialize the ingester with paths.
        
        Args:
            documents_dir: Path to the raw documents.
            retriever: Optional KnowledgeRetriever instance.
        """
        base_dir = Path(__file__).parent.parent
        self.documents_dir = Path(documents_dir) if documents_dir else Path(settings.rag_documents_dir)
        self.retriever = retriever or KnowledgeRetriever()
        
        # Ensure documents directory exists
        self.documents_dir.mkdir(parents=True, exist_ok=True)
        logger.info("Initialized DocumentIngester with directory: %s", self.documents_dir)

    def parse_txt(self, file_path: Path) -> str:
        """Parse TXT file.
        
        Args:
            file_path: Absolute path to the file.
            
        Returns:
            Extracted text content.
        """
        try:
            return file_path.read_text(encoding="utf-8")
        except Exception as e:
            logger.error("Failed to read TXT file %s: %s", file_path, e)
            raise IngestError(f"TXT parse error: {e}") from e

    def parse_md(self, file_path: Path) -> str:
        """Parse Markdown file.
        
        Args:
            file_path: Absolute path to the file.
            
        Returns:
            Extracted text content.
        """
        try:
            return file_path.read_text(encoding="utf-8")
        except Exception as e:
            logger.error("Failed to read Markdown file %s: %s", file_path, e)
            raise IngestError(f"Markdown parse error: {e}") from e

    def parse_docx(self, file_path: Path) -> str:
        """Parse DOCX file using python-docx.
        
        Args:
            file_path: Absolute path to the file.
            
        Returns:
            Extracted text content.
        """
        try:
            import docx
            doc = docx.Document(str(file_path))
            full_text = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                style_name = para.style.name if para.style else ""
                if style_name.startswith("Heading 1"):
                    full_text.append(f"# {text}")
                elif style_name.startswith("Heading 2"):
                    full_text.append(f"## {text}")
                elif style_name.startswith("Heading 3"):
                    full_text.append(f"### {text}")
                else:
                    full_text.append(text)
            return "\n\n".join(full_text)
        except Exception as e:
            logger.error("Failed to read DOCX file %s: %s", file_path, e)
            raise IngestError(f"DOCX parse error: {e}") from e

    def parse_pdf(self, file_path: Path) -> str:
        """Parse PDF file using pypdf, falling back to pdfplumber if available.
        
        Args:
            file_path: Absolute path to the file.
            
        Returns:
            Extracted text content.
        """
        try:
            import pypdf
            reader = pypdf.PdfReader(str(file_path))
            text_parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            return "\n".join(text_parts)
        except ImportError:
            logger.warning("pypdf not installed. Trying pdfplumber...")
            try:
                import pdfplumber  # type: ignore
                with pdfplumber.open(str(file_path)) as pdf:
                    text_parts = [page.extract_text() for page in pdf.pages if page.extract_text()]
                return "\n".join(text_parts)
            except Exception as e:
                logger.error("Could not parse PDF file %s. All PDF parsing libraries failed or missing: %s", file_path, e)
                raise IngestError(f"PDF parsing not supported: libraries missing. Error: {e}") from e
        except Exception as e:
            logger.error("Failed to parse PDF file %s: %s", file_path, e)
            raise IngestError(f"PDF parse error: {e}") from e

    def ingest_file(
        self, 
        file_path: Path, 
        chunk_size: int | None = None, 
        chunk_overlap: int | None = None,
        company_id: str | None = None,
        project_id: str | None = None,
        document_type: str | None = None
    ) -> int:
        """Parse, split, and add a single file to the vector store.
        
        Args:
            file_path: Path to the target file.
            chunk_size: Optional custom chunk size.
            chunk_overlap: Optional custom chunk overlap.
            company_id: Tenant identifier.
            project_id: Project identifier.
            document_type: Type of document (e.g., BRD).
            
        Returns:
            Number of chunks successfully indexed.
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise IngestError(f"File not found: {file_path}")
            
        suffix = file_path.suffix.lower()
        logger.info("Ingesting file: %s (Type: %s)", file_path.name, suffix)
        
        # 1. Parse based on file type
        if suffix == ".txt":
            content = self.parse_txt(file_path)
        elif suffix in (".md", ".markdown"):
            content = self.parse_md(file_path)
        elif suffix == ".docx":
            content = self.parse_docx(file_path)
        elif suffix == ".pdf":
            content = self.parse_pdf(file_path)
        else:
            raise IngestError(f"Unsupported file type: {suffix}")
            
        if not content.strip():
            logger.warning("Empty content extracted from file: %s", file_path)
            return 0
            
        # 2. Chunk contents
        c_size = chunk_size or settings.rag_chunk_size
        c_overlap = chunk_overlap or settings.rag_chunk_overlap
        
        try:
            from langchain_text_splitters import RecursiveCharacterTextSplitter, MarkdownHeaderTextSplitter
            from langchain_core.documents import Document
            import uuid
            
            headers_to_split_on = [
                ("#", "Header 1"),
                ("##", "Header 2"),
                ("###", "Header 3"),
            ]
            
            if suffix in (".md", ".markdown", ".docx"):
                markdown_splitter = MarkdownHeaderTextSplitter(headers_to_split_on=headers_to_split_on, strip_headers=False)
                md_docs = markdown_splitter.split_text(content)
            else:
                md_docs = [Document(page_content=content, metadata={})]
                
            splitter = RecursiveCharacterTextSplitter(chunk_size=c_size, chunk_overlap=c_overlap)
            
            final_chunks = []
            final_metadatas = []
            final_ids = []
            
            global_idx = 0
            
            for doc in md_docs:
                section_name = doc.metadata.get("Header 3") or doc.metadata.get("Header 2") or doc.metadata.get("Header 1") or "General"
                
                # Create Parent chunk
                parent_id = str(uuid.uuid4())
                parent_content = doc.page_content
                
                # Base metadata
                base_meta = {
                    "source": file_path.name,
                    "file_type": suffix,
                    "section": section_name
                }
                if company_id: base_meta["company_id"] = company_id
                if project_id: base_meta["project_id"] = project_id
                if document_type: base_meta["document_type"] = document_type
                
                # Store parent
                parent_meta = base_meta.copy()
                parent_meta["is_parent"] = True
                parent_meta["chunk_type"] = "parent"
                
                final_chunks.append(parent_content)
                final_metadatas.append(parent_meta)
                final_ids.append(parent_id)
                
                # Split parent into children
                child_splits = splitter.split_text(parent_content)
                
                # Store children
                for idx, split in enumerate(child_splits):
                    child_meta = base_meta.copy()
                    child_meta["is_parent"] = False
                    child_meta["chunk_type"] = "child"
                    child_meta["parent_id"] = parent_id
                    child_meta["chunk_index"] = idx
                    
                    prefix = f"{company_id}_{project_id}_" if company_id and project_id else ""
                    child_id = f"{prefix}{file_path.name}_child_{global_idx}"
                    global_idx += 1
                    
                    final_chunks.append(split)
                    final_metadatas.append(child_meta)
                    final_ids.append(child_id)
                    
            chunks = final_chunks
            metadatas = final_metadatas
            ids = final_ids
            
        except ImportError:
            logger.warning("langchain-text-splitters not available. Using local recursive text splitter fallback.")
            chunks = self._fallback_split_text(content, c_size, c_overlap)
            
            metadatas = []
            ids = []
            for idx, chunk in enumerate(chunks):
                meta = {
                    "source": file_path.name,
                    "file_type": suffix,
                    "chunk_index": idx,
                    "total_chunks": len(chunks),
                    "is_parent": False,
                    "chunk_type": "child"
                }
                if company_id: meta["company_id"] = company_id
                if project_id: meta["project_id"] = project_id
                if document_type: meta["document_type"] = document_type
                    
                metadatas.append(meta)
                prefix = f"{company_id}_{project_id}_" if company_id and project_id else ""
                ids.append(f"{prefix}{file_path.name}_chunk_{idx}")

        if not chunks:
            logger.warning("No chunks generated for file: %s", file_path)
            return 0
            
        # --- Priority 2: Ingestion Validation ---
        def validate_ingestion():
            import re
            expected_sections = set()
            for line in content.split('\n'):
                line = line.strip()
                if line.startswith('# ') or line.startswith('## ') or line.startswith('### '):
                    expected_sections.add(line.lstrip('#').strip())
            
            generated_sections = set()
            parents = 0
            children = 0
            for meta in metadatas:
                sec = meta.get("section")
                if sec and sec != "General":
                    generated_sections.add(sec)
                if meta.get("chunk_type") == "parent":
                    parents += 1
                elif meta.get("chunk_type") == "child":
                    children += 1
                    
            missing = expected_sections - generated_sections
            if missing:
                logger.error("Ingestion Validation Failed for %s. Missing sections: %s", file_path.name, missing)
            else:
                logger.info("Ingestion Validation Passed for %s. Detected %d sections, %d parents, %d children.", 
                            file_path.name, len(expected_sections), parents, children)
                            
        validate_ingestion()
        # ----------------------------------------
            
        # 4. Insert chunks
        if self.retriever.use_chroma and self.retriever.collection is not None:
            try:
                self.retriever.collection.upsert(
                    documents=chunks,
                    metadatas=metadatas,
                    ids=ids
                )
            except Exception as e:
                logger.error("Failed to upsert to ChromaDB: %s. Using memory store fallback.", e)
                if self.retriever.memory_store is not None:
                    self.retriever.memory_store.add_texts(chunks, metadatas=metadatas, ids=ids)
        else:
            if self.retriever.memory_store is not None:
                self.retriever.memory_store.add_texts(chunks, metadatas=metadatas, ids=ids)
                
        logger.info("Successfully ingested %s: indexed %d chunks", file_path.name, len(chunks))
        return len(chunks)

    def _fallback_split_text(self, text: str, chunk_size: int, chunk_overlap: int) -> List[str]:
        """A simple implementation of recursive-like character text splitting.
        
        Used when langchain-text-splitters is not available.
        """
        separators = ["\n\n", "\n", " ", ""]
        
        def _split(text_to_split: str, seps: List[str]) -> List[str]:
            if len(text_to_split) <= chunk_size:
                return [text_to_split]
            if not seps:
                return [text_to_split[i:i + chunk_size] for i in range(0, len(text_to_split), chunk_size)]
                
            sep = seps[0]
            parts = text_to_split.split(sep)
            chunks = []
            current_chunk = ""
            
            for part in parts:
                potential_chunk = current_chunk + (sep if current_chunk else "") + part
                if len(potential_chunk) <= chunk_size:
                    current_chunk = potential_chunk
                else:
                    if current_chunk:
                        chunks.append(current_chunk)
                    if len(part) > chunk_size:
                        chunks.extend(_split(part, seps[1:]))
                        current_chunk = ""
                    else:
                        current_chunk = part
            if current_chunk:
                chunks.append(current_chunk)
            return chunks

        return _split(text, separators)

    def process_documents(self) -> dict[str, Any]:
        """Scan the documents directory and ingest all supported files.
        
        Returns:
            Dictionary summarizing the results.
        """
        return self.ingest_directory()

    def ingest_directory(self) -> dict[str, Any]:
        """Scan the documents directory and ingest all supported files.
        
        Returns:
            Dictionary summarizing the results.
        """
        logger.info("Scanning documents directory for ingestion: %s", self.documents_dir)
        if not self.documents_dir.exists():
            logger.warning("Documents directory %s does not exist. Creating it.", self.documents_dir)
            self.documents_dir.mkdir(parents=True, exist_ok=True)
            return {
                "status": "skipped",
                "reason": "documents directory was missing",
                "processed_files": [],
                "failed_files": {},
                "chunks_processed": 0
            }
            
        supported_suffixes = {".txt", ".md", ".markdown", ".docx", ".pdf"}
        files_to_process = [
            f for f in self.documents_dir.iterdir()
            if f.is_file() and f.suffix.lower() in supported_suffixes
        ]
        
        total_chunks = 0
        processed_files = []
        failed_files = {}
        
        for file_path in files_to_process:
            try:
                chunks_created = self.ingest_file(file_path)
                total_chunks += chunks_created
                processed_files.append(file_path.name)
            except Exception as e:
                logger.error("Failed to ingest %s: %s", file_path.name, e)
                failed_files[file_path.name] = str(e)
                
        status = "success" if not failed_files else "partial_success"
        if not processed_files and failed_files:
            status = "failed"
        elif not processed_files and not failed_files:
            status = "success"
            
        return {
            "status": status,
            "processed_files": processed_files,
            "failed_files": failed_files,
            "chunks_processed": total_chunks
        }
